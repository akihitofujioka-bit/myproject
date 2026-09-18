"""core.py の自動テスト。

実行:  python -m unittest discover -s gikai_simple/tests
（gikai_simple フォルダの中から  python -m unittest tests.test_core  でもよい）
"""

from __future__ import annotations

import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent))

import core  # noqa: E402

try:
    from PIL import Image
    PIL_OK = True
except ImportError:
    PIL_OK = False


def make_photo(path: Path, w: int = 1600, h: int = 1200, color=(200, 80, 80)) -> None:
    Image.new("RGB", (w, h), color).save(path, "JPEG")


class ParseTest(unittest.TestCase):
    def test_photo_line_full(self):
        ref = core.parse_photo_line("【写真】森下.jpg｜顔｜森下けい子議員")
        self.assertEqual((ref.file, ref.size, ref.caption), ("森下.jpg", "顔", "森下けい子議員"))
        self.assertEqual(ref.width_mm, core.SIZES["顔"])

    def test_photo_line_half_width_bar_and_defaults(self):
        ref = core.parse_photo_line("  【写真】 a b.png | 大 ")
        self.assertEqual((ref.file, ref.size, ref.caption), ("a b.png", "大", ""))
        ref = core.parse_photo_line("【写真】only.jpg")
        self.assertEqual((ref.file, ref.size, ref.caption), ("only.jpg", "中", ""))
        ref = core.parse_photo_line("【写真】x.jpg｜｜説明だけ")
        self.assertEqual((ref.size, ref.caption), ("中", "説明だけ"))

    def test_not_photo_line(self):
        self.assertIsNone(core.parse_photo_line("写真に兄　樹紀さん"))
        self.assertIsNone(core.parse_photo_line("【一般会計】"))

    def test_blocks_and_count(self):
        text = "見出し\n本文１\n【写真】p.jpg｜小｜説明\n本文２\n"
        blocks = core.parse_blocks(text)
        self.assertEqual([k for k, _ in blocks], ["text", "photo", "text"])
        self.assertEqual(blocks[0][1], "見出し\n本文１")
        self.assertEqual(core.count_chars(text), 9)
        self.assertEqual(core.PhotoRef("p.jpg", "小", "説明").to_line(), "【写真】p.jpg｜小｜説明")


class IssueTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def test_create_and_open(self):
        issue = core.Issue.create(self.root, "２０４", "令和８年７月31日")
        self.assertEqual(issue.folder.name, "第204号")
        self.assertTrue((issue.folder / "01_表紙.txt").exists())
        self.assertTrue((issue.folder / "07_裏表紙.txt").exists())
        self.assertIn("第204号", issue.read_text("表紙"))
        again = core.Issue.open(issue.folder)
        self.assertEqual((again.gou, again.hakkoubi), ("204", "令和８年７月31日"))
        with self.assertRaises(FileExistsError):
            core.Issue.create(self.root, "204", "")

    def test_text_roundtrip_bom(self):
        issue = core.Issue.create(self.root, "1", "")
        issue.write_text("一般質問", "質問　あ\r\n答弁　い")
        raw = issue.text_path("一般質問").read_bytes()
        self.assertTrue(raw.startswith(b"\xef\xbb\xbf"))
        self.assertEqual(issue.read_text("一般質問"), "質問　あ\n答弁　い")

    def test_read_cp932(self):
        p = self.root / "x.txt"
        p.write_bytes("議会だより".encode("cp932"))
        self.assertEqual(core.read_text_file(p), "議会だより")

    def test_import_docx_with_textbox(self):
        p = self.root / "m.docx"
        xml = (
            '<?xml version="1.0"?><w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            "<w:body><w:p><w:r><w:t>本文１</w:t></w:r></w:p>"
            "<w:p><w:r><w:t>本文</w:t></w:r><w:r><w:t>２</w:t></w:r></w:p>"
            "<w:p><w:r><w:pict><w:txbxContent><w:p><w:r><w:t>枠の中</w:t></w:r></w:p></w:txbxContent></w:pict></w:r></w:p>"
            "</w:body></w:document>"
        )
        with zipfile.ZipFile(p, "w") as z:
            z.writestr("word/document.xml", xml)
        self.assertEqual(core.import_manuscript(p), "本文１\n本文２\n枠の中")

    def test_import_docx_rejects_doctype(self):
        p = self.root / "bad.docx"
        with zipfile.ZipFile(p, "w") as z:
            z.writestr("word/document.xml", '<!DOCTYPE x [<!ENTITY a "b">]><w:document/>')
        with self.assertRaises(ValueError):
            core.import_manuscript(p)

    def test_import_unknown_ext(self):
        p = self.root / "a.xlsx"
        p.write_bytes(b"")
        with self.assertRaises(ValueError):
            core.import_manuscript(p)


@unittest.skipUnless(PIL_OK and core.DOCX_OK, "Pillow と python-docx が必要")
class BuildTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.issue = core.Issue.create(Path(self.tmp.name), "204", "令和８年７月31日")
        make_photo(self.issue.photo_dir / "村長.jpg", 1600, 1200)
        make_photo(self.issue.photo_dir / "小さい.jpg", 300, 200)   # 解像度不足
        make_photo(self.issue.photo_dir / "余り.jpg", 800, 600)    # 使わない
        self.issue.write_text("行政報告", "行政報告（要旨）\n【写真】村長.jpg｜顔｜松岡村長\n本文です。")
        self.issue.write_text("特集", "特集の本文\n【写真】小さい.jpg｜大｜広い写真\n【写真】ない.jpg｜中")

    def tearDown(self):
        self.tmp.cleanup()

    def test_usage_and_refs(self):
        usage = self.issue.photo_usage()
        self.assertEqual(usage["村長.jpg"], ["行政報告"])
        self.assertEqual([r.file for _, r in self.issue.all_refs()], ["村長.jpg", "小さい.jpg", "ない.jpg"])

    def test_photo_info(self):
        info = core.photo_info(self.issue.photo_dir / "小さい.jpg")
        self.assertEqual((info.width_px, info.height_px), (300, 200))
        self.assertIn("解像度不足", info.warning_for(80))
        self.assertEqual(info.warning_for(20), "")
        self.assertIn("ありません", core.photo_info(self.issue.photo_dir / "ない.jpg").error)

    def test_build_all(self):
        outs, warnings = core.build_all(self.issue)
        self.assertEqual(len(outs), 2)
        for p in outs:
            self.assertTrue(p.exists() and p.stat().st_size > 0, p)
        joined = "\n".join(warnings)
        self.assertIn("ない.jpg", joined)
        self.assertIn("解像度不足", joined)
        self.assertNotIn("表紙: 原稿が空です", warnings)   # 表紙は雛形入りなので空ではない
        self.assertIn("一般質問: 原稿が空です", warnings)

        from docx import Document
        doc = Document(outs[0])
        texts = [p.text for p in doc.paragraphs]
        self.assertTrue(any(t.startswith("■ 行政報告") for t in texts))
        self.assertTrue(any(t.startswith("【写真1】村長.jpg（顔・幅26mm）") and "松岡村長" in t for t in texts))
        self.assertTrue(any(t.startswith("【写真3】ない.jpg") for t in texts))
        self.assertEqual(len(doc.inline_shapes), 2)   # 見つかった写真 2 枚だけ貼られる
        # 原稿は縦書き、指示書は横書き
        from docx.oxml.ns import qn
        td = doc.sections[0]._sectPr.find(qn("w:textDirection"))
        self.assertIsNotNone(td)
        self.assertEqual(td.get(qn("w:val")), "tbRl")
        self.assertEqual(list(doc.sections[0]._sectPr).index(td) + 1,
                         list(doc.sections[0]._sectPr).index(doc.sections[0]._sectPr.find(qn("w:docGrid"))))

        sheet = Document(outs[1])
        self.assertIsNone(sheet.sections[0]._sectPr.find(qn("w:textDirection")))
        table = sheet.tables[0]
        self.assertEqual(len(table.rows), 1 + 3)
        self.assertEqual(table.rows[1].cells[2].paragraphs[0].text, "村長.jpg")
        self.assertIn("余り.jpg", "\n".join(p.text for p in sheet.paragraphs))


if __name__ == "__main__":
    unittest.main()
