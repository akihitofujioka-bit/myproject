"""議会だより かんたん原稿ツール — 中核（画面なしで動く部分）

1号ぶんを「1つのフォルダ」で扱う。

    第204号/
      号情報.json                 号数・発行日
      01_表紙.txt                 区分ごとの原稿（この中に【写真】行を書く）
      02_行政報告.txt
      ...
      07_裏表紙.txt
      写真/                       写真の原本（印刷所にそのまま渡す）
      出力/                       作った Word

写真の配置は、原稿の中に次の 1 行を書くだけで表す。

    【写真】ファイル名.jpg｜中｜説明文

区切りは全角「｜」でも半角「|」でもよい。大きさは 大・中・小・顔 の 4 つ。
大きさと説明は省略できる。ツールは、この行を Word 原稿ではその場所に
縮小した写真として貼り、写真配置指示書では一覧にまとめる。

推測は一切しない。書いてあるとおりに並べるだけ。
"""

from __future__ import annotations

import datetime
import io
import json
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from xml.etree import ElementTree as ET

try:
    from PIL import Image, ImageOps
    PIL_OK = True
except ImportError:  # 写真の縮小・向き補正ができないだけで、他は動く
    PIL_OK = False

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt, RGBColor
    DOCX_OK = True
except ImportError:
    DOCX_OK = False

import doc97

# ---------------------------------------------------------------- 決まりごと

# 紙面の順。番号がファイル名の頭に付く。
KUBUN = [
    ("01", "表紙"),
    ("02", "行政報告"),
    ("03", "審議したこと・決まったこと"),
    ("04", "閉会中の委員会活動報告"),
    ("05", "一般質問"),
    ("06", "特集"),
    ("07", "裏表紙"),
]

# 写真の大きさ → 幅（mm）。紙面は 5 段組で 1 段が約 30mm。
SIZES = {
    "大": 80,   # 3 段ぶん弱（段の見出し写真など）
    "中": 55,   # 2 段ぶん
    "小": 38,   # 1 段と少し
    "顔": 26,   # 議員・委員長の顔写真
}
DEFAULT_SIZE = "中"

PHOTO_EXT = {".jpg", ".jpeg", ".png", ".gif", ".bmp", ".tif", ".tiff", ".webp", ".heic"}
PRINT_DPI = 350   # 印刷所が求める解像度

INFO_NAME = "号情報.json"
PHOTO_DIR = "写真"
OUT_DIR = "出力"

FONT_MINCHO = "ＭＳ 明朝"
FONT_GOTHIC = "ＭＳ ゴシック"

# 【写真】ファイル名｜大きさ｜説明
PHOTO_LINE = re.compile(
    r"^\s*【写真】\s*(?P<file>[^｜|]+?)"
    r"(?:\s*[｜|]\s*(?P<size>[大中小顔])?)?"
    r"(?:\s*[｜|]\s*(?P<caption>.*?))?\s*$"
)

# 表紙の原稿に最初から入れておく文面
COVER_TEMPLATE = """第{gou}号
{hakkoubi}
ひだか議会だより

特集　（特集の題名）　…………（ページ）

発行　高知県日高村議会　編集　議会広報発行調査特別委員会
日高村本郷６１－１　〒７８１－２１９４　℡０８８９－２４－７７７７
"""

READ_ME = """このフォルダの使い方

・01_表紙.txt 〜 07_裏表紙.txt に、区分ごとの原稿を書きます（紙面の順です）
・写真 フォルダに写真の原本を入れます（ファイル名はそのまま印刷所に伝わります）
・写真を入れたい場所に、次の 1 行を書きます

    【写真】ファイル名.jpg｜中｜説明文

  大きさは 大・中・小・顔 のどれか（省略すると 中）。説明文は省略できます。
・ツールで「Word を作る」を押すと、出力 フォルダに
    原稿（写真入り）と 写真配置指示書 の 2 つの Word ができます
・印刷所には「出力の Word 2 つ」と「写真フォルダの中身」を渡します
"""


# ---------------------------------------------------------------- 原稿の解釈

@dataclass
class PhotoRef:
    """原稿の中の【写真】行 1 つ。"""
    file: str
    size: str = DEFAULT_SIZE
    caption: str = ""
    line_no: int = 0

    @property
    def width_mm(self) -> int:
        return SIZES.get(self.size, SIZES[DEFAULT_SIZE])

    def to_line(self) -> str:
        line = f"【写真】{self.file}｜{self.size}"
        if self.caption:
            line += f"｜{self.caption}"
        return line


def parse_photo_line(line: str, line_no: int = 0) -> PhotoRef | None:
    m = PHOTO_LINE.match(line)
    if not m:
        return None
    return PhotoRef(
        file=m.group("file").strip(),
        size=(m.group("size") or DEFAULT_SIZE),
        caption=(m.group("caption") or "").strip(),
        line_no=line_no,
    )


def parse_blocks(text: str) -> list[tuple[str, object]]:
    """原稿を、("text", 文字列) と ("photo", PhotoRef) の並びに分ける。

    連続する本文行は 1 つの "text" にまとめる。空行はそのまま残す。
    """
    blocks: list[tuple[str, object]] = []
    buf: list[str] = []
    for i, line in enumerate(text.splitlines(), 1):
        ref = parse_photo_line(line, i)
        if ref:
            if buf:
                blocks.append(("text", "\n".join(buf)))
                buf = []
            blocks.append(("photo", ref))
        else:
            buf.append(line)
    if buf:
        blocks.append(("text", "\n".join(buf)))
    return blocks


def photo_refs(text: str) -> list[PhotoRef]:
    return [b for kind, b in parse_blocks(text) if kind == "photo"]


def count_chars(text: str) -> int:
    """本文の字数（【写真】行と空白を除く）。枠に入るかの目安に使う。"""
    n = 0
    for line in text.splitlines():
        if parse_photo_line(line):
            continue
        n += len(re.sub(r"\s", "", line))
    return n


# ---------------------------------------------------------------- 号フォルダ

@dataclass
class Issue:
    """1 号ぶんのフォルダ。"""
    folder: Path
    gou: str = ""
    hakkoubi: str = ""

    # ---- 作る・開く

    @classmethod
    def create(cls, root: Path | str, gou: str, hakkoubi: str) -> "Issue":
        """root の下に「第○号」フォルダを作り、空の原稿ファイルをそろえる。"""
        gou = str(gou).strip().translate(str.maketrans("０１２３４５６７８９", "0123456789"))
        folder = Path(root) / f"第{gou}号"
        if folder.exists():
            raise FileExistsError(f"{folder} は既にあります")
        folder.mkdir(parents=True)
        (folder / PHOTO_DIR).mkdir()
        (folder / OUT_DIR).mkdir()
        issue = cls(folder=folder, gou=gou, hakkoubi=hakkoubi)
        issue.save_info()
        for code, name in KUBUN:
            body = COVER_TEMPLATE.format(gou=gou, hakkoubi=hakkoubi) if name == "表紙" else ""
            issue.write_text(name, body)
        (folder / "はじめにお読みください.txt").write_text(READ_ME, encoding="utf-8-sig")
        return issue

    @classmethod
    def open(cls, folder: Path | str) -> "Issue":
        folder = Path(folder)
        if not folder.is_dir():
            raise FileNotFoundError(f"{folder} が見つかりません")
        issue = cls(folder=folder)
        p = folder / INFO_NAME
        if p.exists():
            info = json.loads(p.read_text(encoding="utf-8-sig"))
            issue.gou = str(info.get("gou", ""))
            issue.hakkoubi = str(info.get("hakkoubi", ""))
        else:
            m = re.search(r"(\d+)", folder.name)
            issue.gou = m.group(1) if m else ""
        (folder / PHOTO_DIR).mkdir(exist_ok=True)
        (folder / OUT_DIR).mkdir(exist_ok=True)
        return issue

    def save_info(self) -> None:
        (self.folder / INFO_NAME).write_text(
            json.dumps({"gou": self.gou, "hakkoubi": self.hakkoubi}, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )

    # ---- 原稿

    def text_path(self, kubun: str) -> Path:
        for code, name in KUBUN:
            if name == kubun:
                return self.folder / f"{code}_{name}.txt"
        raise KeyError(kubun)

    def read_text(self, kubun: str) -> str:
        p = self.text_path(kubun)
        return read_text_file(p) if p.exists() else ""

    def write_text(self, kubun: str, text: str) -> None:
        # BOM 付き UTF-8 にしておくと、Windows のメモ帳でも文字化けしない
        self.text_path(kubun).write_text(text.replace("\r\n", "\n"), encoding="utf-8-sig")

    def all_texts(self) -> list[tuple[str, str]]:
        return [(name, self.read_text(name)) for _, name in KUBUN]

    # ---- 写真

    @property
    def photo_dir(self) -> Path:
        return self.folder / PHOTO_DIR

    @property
    def out_dir(self) -> Path:
        return self.folder / OUT_DIR

    def photo_files(self) -> list[Path]:
        if not self.photo_dir.is_dir():
            return []
        files = [p for p in self.photo_dir.iterdir()
                 if p.is_file() and p.suffix.lower() in PHOTO_EXT and not p.name.startswith(".")]
        return sorted(files, key=lambda p: natural_key(p.name))

    def photo_usage(self) -> dict[str, list[str]]:
        """写真ファイル名 → 使われている区分の一覧。"""
        used: dict[str, list[str]] = {}
        for name, text in self.all_texts():
            for ref in photo_refs(text):
                used.setdefault(ref.file, []).append(name)
        return used

    def all_refs(self) -> list[tuple[str, PhotoRef]]:
        """紙面の順に (区分, PhotoRef) を並べる。"""
        out = []
        for name, text in self.all_texts():
            for ref in photo_refs(text):
                out.append((name, ref))
        return out

    def title(self) -> str:
        return f"第{self.gou}号" if self.gou else self.folder.name


def natural_key(s: str):
    return [int(t) if t.isdigit() else t for t in re.split(r"(\d+)", s)]


# ---------------------------------------------------------------- 取り込み

def read_text_file(path: Path | str) -> str:
    """文字コードを自動判別してテキストを読む（UTF-8 → CP932 → EUC-JP）。"""
    raw = Path(path).read_bytes()
    for enc in ("utf-8-sig", "cp932", "euc_jp"):
        try:
            return raw.decode(enc).replace("\r\n", "\n").replace("\r", "\n")
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def import_manuscript(path: Path | str) -> str:
    """議員から届いた原稿ファイルから文字だけを取り出す。

    .docx（本文・表・テキストボックス）、.doc（旧形式）、.txt に対応。
    書式は捨てる。段落の区切りだけ改行として残す。
    """
    path = Path(path)
    ext = path.suffix.lower()
    if ext == ".docx":
        return docx_to_text(path)
    if ext == ".doc":
        return doc97.extract_text(path)
    if ext in (".txt", ".text", ".md"):
        return read_text_file(path)
    raise ValueError(f"この形式は取り込めません: {path.name}（.docx / .doc / .txt に対応）")


_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def docx_to_text(path: Path | str) -> str:
    """python-docx を使わず、document.xml から段落ごとの文字を取り出す。

    テキストボックス（txbxContent）の中も本文と同じ順に拾う。
    """
    with zipfile.ZipFile(path) as z:
        xml = z.read("word/document.xml")
    # Word が作る XML に DOCTYPE や実体定義は含まれない。含まれていたら
    # 細工されたファイルなので読まない（XML 爆弾・外部実体参照の対策）
    head = xml[:4096].lower()
    if b"<!doctype" in head or b"<!entity" in xml.lower():
        raise ValueError(f"通常の Word ファイルではありません: {Path(path).name}")
    root = ET.fromstring(xml)
    lines: list[str] = []
    for p in root.iter(_W + "p"):
        parts: list[str] = []
        has_nested = False

        def walk(node):
            nonlocal has_nested
            for ch in node:
                if ch.tag == _W + "p":      # テキストボックス内の段落は別に拾う
                    has_nested = True
                    continue
                if ch.tag == _W + "t":
                    parts.append(ch.text or "")
                elif ch.tag in (_W + "br", _W + "cr"):
                    parts.append("\n")
                elif ch.tag == _W + "tab":
                    parts.append("\t")
                else:
                    walk(ch)

        walk(p)
        line = "".join(parts)
        if not line and has_nested:   # 枠だけの段落は空行にしない
            continue
        lines.append(line)
    text = "\n".join(lines)
    return re.sub(r"\n{3,}", "\n\n", text).strip("\n")


# ---------------------------------------------------------------- 写真の情報

@dataclass
class PhotoInfo:
    path: Path
    width_px: int = 0
    height_px: int = 0
    error: str = ""

    def max_print_mm(self) -> float:
        """350dpi で刷ったとき、何 mm 幅まで使えるか。"""
        return self.width_px / PRINT_DPI * 25.4 if self.width_px else 0.0

    def warning_for(self, width_mm: int) -> str:
        if self.error:
            return self.error
        limit = self.max_print_mm()
        if limit and limit < width_mm:
            return f"解像度不足（{width_mm}mm には {limit:.0f}mm 相当までしか使えません。元データをもらってください）"
        return ""


def photo_info(path: Path | str) -> PhotoInfo:
    path = Path(path)
    if not path.exists():
        return PhotoInfo(path=path, error="写真フォルダにありません")
    if not PIL_OK:
        return PhotoInfo(path=path)
    try:
        with Image.open(path) as im:
            im = ImageOps.exif_transpose(im)
            return PhotoInfo(path=path, width_px=im.width, height_px=im.height)
    except Exception as e:  # 壊れた画像など
        return PhotoInfo(path=path, error=f"読み込めません（{e}）")


def reduced_jpeg(path: Path | str, max_px: int = 1200) -> io.BytesIO | None:
    """Word に貼る用に、向きを直して縮小した JPEG をメモリ上に作る。

    原本は触らない。Word のファイルが重くならないようにするため。
    """
    if not PIL_OK:
        return None
    try:
        with Image.open(path) as im:
            im = ImageOps.exif_transpose(im)
            im.thumbnail((max_px, max_px))
            if im.mode not in ("RGB", "L"):
                im = im.convert("RGB")
            buf = io.BytesIO()
            im.save(buf, "JPEG", quality=85)
            buf.seek(0)
            return buf
    except Exception:
        return None


# ---------------------------------------------------------------- Word 生成

def _set_font(run, name=FONT_MINCHO, size=10.5, bold=False, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), name)


def _para(doc, text="", *, name=FONT_MINCHO, size=10.5, bold=False,
          color=None, align=None, after=4):
    p = doc.add_paragraph()
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(after)
    if text:
        _set_font(p.add_run(text), name, size, bold, color)
    return p


def _setup_page(doc):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.top_margin = sec.bottom_margin = Mm(20)
    sec.left_margin = sec.right_margin = Mm(20)
    style = doc.styles["Normal"]
    style.font.name = FONT_MINCHO
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_MINCHO)


def _add_photo(doc, issue: Issue, number: int, kubun: str, ref: PhotoRef,
               warnings: list[str]) -> None:
    """原稿 Word の中に、写真（縮小版）と赤い指示文を入れる。"""
    src = issue.photo_dir / ref.file
    info = photo_info(src)
    label = f"【写真{number}】{ref.file}（{ref.size}・幅{ref.width_mm}mm）"
    if ref.caption:
        label += f"　説明: {ref.caption}"
    _para(doc, label, name=FONT_GOTHIC, size=9, bold=True, color=(0xC0, 0x00, 0x00), after=1)

    if info.error:
        warnings.append(f"{kubun}: {ref.file} — {info.error}")
        _para(doc, f"（{info.error}）", name=FONT_GOTHIC, size=9, color=(0xC0, 0x00, 0x00))
        return
    warn = info.warning_for(ref.width_mm)
    if warn:
        warnings.append(f"{kubun}: {ref.file} — {warn}")

    buf = reduced_jpeg(src)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    try:
        p.add_run().add_picture(buf if buf else str(src), width=Mm(ref.width_mm))
    except Exception as e:
        warnings.append(f"{kubun}: {ref.file} — Word に貼れませんでした（{e}）")
        _set_font(p.add_run("（写真を貼れませんでした）"), FONT_GOTHIC, 9, color=(0xC0, 0, 0))
    if warn:
        _para(doc, f"※ {warn}", name=FONT_GOTHIC, size=8, color=(0xC0, 0x00, 0x00))


def build_manuscript(issue: Issue, out: Path | str | None = None) -> tuple[Path, list[str]]:
    """原稿 Word（写真入り）を作る。戻り値は (保存先, 注意事項)。"""
    if not DOCX_OK:
        raise RuntimeError("python-docx が入っていません（pip install python-docx）")
    out = Path(out) if out else issue.out_dir / f"{issue.title()}_原稿.docx"
    doc = Document()
    _setup_page(doc)
    warnings: list[str] = []

    _para(doc, f"ひだか議会だより {issue.title()}　原稿", name=FONT_GOTHIC, size=16, bold=True, align="center")
    _para(doc, f"発行日 {issue.hakkoubi}　／　作成 {datetime.date.today():%Y-%m-%d}",
          name=FONT_GOTHIC, size=9, align="center", after=12)
    _para(doc, "赤い【写真○】は写真を入れる場所と大きさの指示です。写真の原本は「写真」フォルダにあります。",
          name=FONT_GOTHIC, size=9, color=(0xC0, 0, 0), after=12)

    number = 0
    for idx, (kubun, text) in enumerate(issue.all_texts()):
        if idx > 0:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        _para(doc, f"■ {kubun}", name=FONT_GOTHIC, size=14, bold=True, after=8)
        if not text.strip():
            _para(doc, "（原稿なし）", name=FONT_GOTHIC, size=9, color=(0x80, 0x80, 0x80))
            warnings.append(f"{kubun}: 原稿が空です")
        for kind, block in parse_blocks(text):
            if kind == "photo":
                number += 1
                _add_photo(doc, issue, number, kubun, block, warnings)
            else:
                for line in block.split("\n"):
                    _para(doc, line, after=0)
    doc.save(out)
    return out, warnings


def build_photo_sheet(issue: Issue, out: Path | str | None = None) -> tuple[Path, list[str]]:
    """写真配置指示書（一覧表）を作る。印刷所が写真を照合するための紙。"""
    if not DOCX_OK:
        raise RuntimeError("python-docx が入っていません（pip install python-docx）")
    out = Path(out) if out else issue.out_dir / f"{issue.title()}_写真配置指示.docx"
    doc = Document()
    _setup_page(doc)
    warnings: list[str] = []

    _para(doc, f"ひだか議会だより {issue.title()}　写真配置指示書", name=FONT_GOTHIC, size=16, bold=True, align="center")
    _para(doc, f"作成 {datetime.date.today():%Y-%m-%d}　写真の原本は「写真」フォルダ内。番号は原稿 Word の【写真○】と同じ。",
          name=FONT_GOTHIC, size=9, align="center", after=10)

    refs = issue.all_refs()
    if not refs:
        _para(doc, "写真の指定がありません。", name=FONT_GOTHIC)
    else:
        table = doc.add_table(rows=1, cols=6)
        table.style = "Table Grid"
        heads = ["番号", "写真", "ファイル名", "区分", "大きさ", "説明"]
        for cell, h in zip(table.rows[0].cells, heads):
            cell.text = ""
            _set_font(cell.paragraphs[0].add_run(h), FONT_GOTHIC, 9, True)
        for n, (kubun, ref) in enumerate(refs, 1):
            row = table.add_row().cells
            _set_font(row[0].paragraphs[0].add_run(str(n)), FONT_GOTHIC, 10, True)
            src = issue.photo_dir / ref.file
            info = photo_info(src)
            if info.error:
                _set_font(row[1].paragraphs[0].add_run(info.error), FONT_GOTHIC, 8, color=(0xC0, 0, 0))
                warnings.append(f"{kubun}: {ref.file} — {info.error}")
            else:
                buf = reduced_jpeg(src, 600)
                try:
                    row[1].paragraphs[0].add_run().add_picture(buf if buf else str(src), width=Mm(35))
                except Exception:
                    _set_font(row[1].paragraphs[0].add_run("（表示不可）"), FONT_GOTHIC, 8)
            _set_font(row[2].paragraphs[0].add_run(ref.file), FONT_GOTHIC, 9)
            if info.width_px:
                _set_font(row[2].add_paragraph().add_run(f"{info.width_px}×{info.height_px}px"), FONT_GOTHIC, 8)
            _set_font(row[3].paragraphs[0].add_run(kubun), FONT_GOTHIC, 9)
            _set_font(row[4].paragraphs[0].add_run(f"{ref.size}（幅{ref.width_mm}mm）"), FONT_GOTHIC, 9)
            _set_font(row[5].paragraphs[0].add_run(ref.caption or "（説明なし）"), FONT_MINCHO, 9)
            warn = info.warning_for(ref.width_mm)
            if warn:
                _set_font(row[5].add_paragraph().add_run("※ " + warn), FONT_GOTHIC, 8, color=(0xC0, 0, 0))
                warnings.append(f"{kubun}: {ref.file} — {warn}")
        widths = [12, 40, 40, 30, 22, 26]
        for row in table.rows:
            for cell, w in zip(row.cells, widths):
                cell.width = Mm(w)

    # 使っていない写真も知らせる（入れ忘れの確認用）
    used = issue.photo_usage()
    unused = [p.name for p in issue.photo_files() if p.name not in used]
    if unused:
        _para(doc, "", after=8)
        _para(doc, "写真フォルダにあるが、原稿で使っていない写真:", name=FONT_GOTHIC, size=9, bold=True)
        for name in unused:
            _para(doc, "・" + name, name=FONT_GOTHIC, size=9, after=0)
    doc.save(out)
    return out, warnings


def build_all(issue: Issue) -> tuple[list[Path], list[str]]:
    """原稿 Word と写真配置指示書をまとめて作る。"""
    outs: list[Path] = []
    warnings: list[str] = []
    for fn in (build_manuscript, build_photo_sheet):
        path, w = fn(issue)
        outs.append(path)
        for item in w:
            if item not in warnings:
                warnings.append(item)
    return outs, warnings
