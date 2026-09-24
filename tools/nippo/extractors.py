"""素材ファイルから文字を取り出す。

対応する種類:
    テキスト   .txt .md .csv .tsv .json .log
    Word       .docx（python-docx）、.doc .rtf .odt（macOS 標準の textutil）
    Excel      .xlsx .xlsm（openpyxl）。旧形式 .xls は未対応
    画像       .jpg .png .heic など → OCR。手書きメモもここ
    PDF        文字層があればそれを使い、なければページを画像にして OCR
    音声・動画 .m4a .mp3 .wav .mov など → Whisper（ローカル）で文字起こし。**macOS のみ**

macOS と Windows のどちらでも動く。OS によって使う道具が変わる:

    | 処理        | macOS                  | Windows                        |
    |-------------|------------------------|--------------------------------|
    | 画像の OCR  | Vision（標準）         | Windows.Media.Ocr（標準）      |
    | PDF の OCR  | PDFKit（標準）         | PyMuPDF でページを画像にして OCR |
    | .doc 等     | textutil（標準）       | .doc は doc97 で読む（.rtf/.odt は未対応） |
    | 撮影日時    | sips（標準）           | Pillow の EXIF                 |
    | 音声        | Whisper                | 未対応（文字起こし済みのテキストを入れてください） |

すべてこのパソコンの中だけで処理し、外部に送らない。
取り出した結果は workspace/.cache に保存し、同じファイルを二度と処理しない。
"""

import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Callable, Dict, Optional

HERE = Path(__file__).resolve().parent
OCR_SOURCE = HERE / "ocr_helper.swift"          # macOS 用（Vision）
OCR_BINARY = HERE / ".build" / "nippo_ocr"
OCR_PS1 = HERE / "ocr_windows.ps1"              # Windows 用（Windows.Media.Ocr）

IS_MAC = sys.platform == "darwin"
IS_WIN = sys.platform == "win32"

# ffmpeg（Whisper が音声の読み込みに使う）が置かれがちな場所を PATH に足す。
# Finder からダブルクリックで起動したときは PATH が最小限になるため。
for extra in ("~/bin", "/opt/homebrew/bin", "/usr/local/bin"):
    p = os.path.expanduser(extra)
    if os.path.isdir(p) and p not in os.environ.get("PATH", "").split(os.pathsep):
        os.environ["PATH"] = p + os.pathsep + os.environ.get("PATH", "")

KINDS = {
    "text": {".txt", ".md", ".markdown", ".csv", ".tsv", ".json", ".log"},
    "word": {".docx", ".doc", ".rtf", ".odt"},
    "excel": {".xlsx", ".xlsm", ".xls"},
    "image": {".jpg", ".jpeg", ".png", ".heic", ".heif", ".tif", ".tiff", ".gif", ".bmp", ".webp"},
    "pdf": {".pdf"},
    "audio": {".m4a", ".mp3", ".wav", ".aac", ".aiff", ".aif", ".flac", ".ogg", ".opus", ".caf",
              ".mp4", ".mov", ".wma", ".webm"},
}
KIND_LABEL = {"text": "テキスト", "word": "Word", "excel": "Excel", "image": "画像", "pdf": "PDF", "audio": "音声"}

Log = Callable[[str], None]


def kind_of(path: Path) -> Optional[str]:
    ext = path.suffix.lower()
    for kind, exts in KINDS.items():
        if ext in exts:
            return kind
    return None


def is_material(path: Path) -> bool:
    """素材として扱うファイルか（隠しファイル・Office の一時ファイルは除く）。"""
    if not path.is_file():
        return False
    if path.name.startswith(".") or path.name.startswith("~$"):
        return False
    return kind_of(path) is not None


# ---------------------------------------------------------------- 日付の判定

_DATE_PATTERNS = [
    re.compile(r"(20\d{2})[-_./年](\d{1,2})[-_./月](\d{1,2})"),
    re.compile(r"(?<!\d)(20\d{2})(\d{2})(\d{2})(?!\d)"),
]
_TIME_PATTERN = re.compile(r"(?<!\d)(\d{1,2})[:時](\d{2})(?!\d)")


def _date_from_name(name: str) -> Optional[datetime]:
    for pat in _DATE_PATTERNS:
        m = pat.search(name)
        if m:
            try:
                d = datetime(int(m.group(1)), int(m.group(2)), int(m.group(3)))
            except ValueError:
                continue
            t = _TIME_PATTERN.search(name[m.end():])
            if t and int(t.group(1)) < 24:
                d = d.replace(hour=int(t.group(1)), minute=int(t.group(2)))
            return d
    return None


def _date_from_exif(path: Path) -> Optional[datetime]:
    """写真の撮影日時。macOS は sips、Windows は Pillow が EXIF から読む。"""
    if IS_MAC:
        try:
            out = subprocess.run(["sips", "-g", "creation", str(path)],
                                 capture_output=True, text=True, timeout=20).stdout
        except Exception:
            return None
        m = re.search(r"creation:\s*(\d{4}):(\d{2}):(\d{2}) (\d{2}):(\d{2}):(\d{2})", out)
        if not m:
            return None
        try:
            return datetime(*map(int, m.groups()))
        except ValueError:
            return None
    # Windows / その他
    try:
        from PIL import Image
        with Image.open(path) as im:
            exif = im.getexif()
        for tag in (36867, 36868, 306):   # DateTimeOriginal, DateTimeDigitized, DateTime
            v = exif.get(tag)
            if v:
                return datetime.strptime(str(v).strip(), "%Y:%m:%d %H:%M:%S")
    except Exception:
        return None
    return None


def _date_from_media(path: Path) -> Optional[datetime]:
    """録音・動画の作成日時（ffmpeg のメタデータ。UTC で入っているので日本時間に直す）。"""
    if not shutil.which("ffmpeg"):
        return None
    try:
        err = subprocess.run(["ffmpeg", "-hide_banner", "-i", str(path)], capture_output=True, text=True, timeout=20).stderr
    except Exception:
        return None
    m = re.search(r"creation_time\s*:\s*(\d{4})-(\d{2})-(\d{2})T(\d{2}):(\d{2}):(\d{2})", err)
    if not m:
        return None
    try:
        utc = datetime(*map(int, m.groups()), tzinfo=timezone.utc)
        return utc.astimezone().replace(tzinfo=None)
    except ValueError:
        return None


def detect_datetime(path: Path, inbox: Path) -> "tuple[datetime, str]":
    """素材の日付を決める。戻り値は (日時, 判定根拠)。"""
    d = _date_from_name(path.stem)
    if d:
        return d, "ファイル名"
    for parent in path.parents:  # 素材フォルダの直下まで、親フォルダ名を順に見る
        if parent == inbox or inbox not in parent.parents:
            break
        d = _date_from_name(parent.name)
        if d:
            return d, "フォルダ名"
    kind = kind_of(path)
    if kind == "image":
        d = _date_from_exif(path)
        if d:
            return d, "撮影日時"
    if kind == "audio":
        d = _date_from_media(path)
        if d:
            return d, "録音日時"
    st = path.stat()
    born = getattr(st, "st_birthtime", None) or st.st_mtime
    return datetime.fromtimestamp(born), "ファイル作成日時"


# ---------------------------------------------------------------- 各形式の読み取り

def _read_text(path: Path) -> str:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "cp932", "euc_jp"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _textutil(path: Path) -> str:
    """.doc / .rtf / .odt から文字を取り出す。

    macOS は標準の textutil が全部読める。Windows にはそれが無いので、
    .doc だけは議会だよりツールの doc97（標準ライブラリだけで旧 Word を読む）を借りる。
    .rtf / .odt は Windows では読めないので、その旨を伝えて終わる。
    """
    if IS_MAC:
        r = subprocess.run(["textutil", "-convert", "txt", "-stdout", str(path)],
                           capture_output=True, timeout=120)
        if r.returncode != 0:
            raise RuntimeError(r.stderr.decode("utf-8", errors="replace").strip() or "textutil が失敗")
        return r.stdout.decode("utf-8", errors="replace")

    ext = path.suffix.lower()
    if ext == ".doc":
        doc97 = _load_doc97()
        if doc97 is None:
            raise RuntimeError(
                "旧形式 .doc を読む部品が見つかりません。"
                "tools/gikai_simple/doc97.py をこのフォルダにコピーしてください")
        try:
            return doc97.extract_text(path)
        except Exception as e:
            raise RuntimeError("この .doc は読めませんでした（%s）。"
                               "Word で開いて .docx として保存し直してください" % e)
    raise RuntimeError(
        "%s は Windows では読めません。Word で開いて .docx として保存し直すか、"
        "内容をテキストファイルに貼り付けてください" % ext)


_doc97_module = None


def _load_doc97():
    """doc97.py を探して読み込む（同じフォルダ → tools/gikai_simple の順）。"""
    global _doc97_module
    if _doc97_module is not None:
        return _doc97_module or None
    import importlib.util
    for cand in (HERE / "doc97.py", HERE.parent / "gikai_simple" / "doc97.py"):
        if cand.exists():
            spec = importlib.util.spec_from_file_location("nippo_doc97", cand)
            mod = importlib.util.module_from_spec(spec)
            spec.loader.exec_module(mod)
            _doc97_module = mod
            return mod
    _doc97_module = False
    return None


def _read_word(path: Path) -> str:
    if path.suffix.lower() == ".docx":
        try:
            import docx  # python-docx
        except ImportError:
            return _textutil(path)
        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs if p.text.strip()]
        for table in d.tables:
            for i, row in enumerate(table.rows):
                cells = [c.text.strip().replace("\n", " ") for c in row.cells]
                if any(cells):  # 表の 1 行目は見出し扱い（本文には載せない）
                    parts.append(("【見出し】" if i == 0 else "") + " | ".join(cells))
        return "\n".join(parts)
    return _textutil(path)


def _read_excel(path: Path) -> str:
    if path.suffix.lower() == ".xls":
        raise RuntimeError("旧形式 .xls は未対応です。Excel で .xlsx として保存し直してください")
    import openpyxl
    wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    parts = []
    for ws in wb.worksheets:
        parts.append("【シート: %s】" % ws.title)
        n = 0
        for row in ws.iter_rows(values_only=True):
            cells = ["" if v is None else str(v).strip() for v in row]
            if not any(cells):
                continue
            # 右端の空セルを落として見やすくする
            while cells and cells[-1] == "":
                cells.pop()
            # 最初の行は見出しとみなす（報告書の本文には載せず、素材欄にだけ残す）
            parts.append(("【見出し】" if n == 0 else "") + " | ".join(cells))
            n += 1
            if n >= 2000:
                parts.append("（2000 行を超えたため以降は省略）")
                break
    return "\n".join(parts)


def ensure_ocr_helper(log: Log) -> Path:
    """OCR 補助プログラムがなければ（または Swift 側が新しければ）swiftc で組み立てる。"""
    if OCR_BINARY.exists() and OCR_BINARY.stat().st_mtime >= OCR_SOURCE.stat().st_mtime:
        return OCR_BINARY
    if not shutil.which("swiftc"):
        raise RuntimeError("swiftc が見つかりません。Xcode（または Command Line Tools）を入れてください")
    log("OCR 補助プログラムを組み立てています（初回のみ、1 分ほど）…")
    OCR_BINARY.parent.mkdir(parents=True, exist_ok=True)
    cmd = ["swiftc", "-O", "-o", str(OCR_BINARY), str(OCR_SOURCE),
           "-framework", "Vision", "-framework", "PDFKit", "-framework", "AppKit"]
    r = subprocess.run(cmd, capture_output=True, text=True, timeout=600)
    if r.returncode != 0 or not OCR_BINARY.exists():
        raise RuntimeError("OCR 補助プログラムの組み立てに失敗:\n" + r.stderr.strip())
    return OCR_BINARY


def _read_ocr(path: Path, log: Log) -> Dict:
    """画像・PDF から文字を取り出す。使う道具は OS で変わる。"""
    if IS_MAC:
        binary = ensure_ocr_helper(log)
        r = subprocess.run([str(binary), str(path)], capture_output=True, text=True, timeout=600)
        if r.returncode != 0:
            raise RuntimeError(r.stderr.strip() or "OCR が失敗")
        data = json.loads(r.stdout.strip().splitlines()[-1])
        return {"text": data.get("text", ""), "confidence": float(data.get("confidence", 0)),
                "pages": int(data.get("pages", 1)), "method": data.get("method", "ocr")}
    if IS_WIN:
        return _read_ocr_windows(path, log)
    raise RuntimeError("この OS では OCR を使えません（macOS か Windows が要ります）")


def _ocr_windows_image(path: Path) -> Dict:
    """Windows 標準の OCR を PowerShell 経由で呼ぶ（画像 1 枚）。"""
    if not OCR_PS1.exists():
        raise RuntimeError("ocr_windows.ps1 が見つかりません")
    r = subprocess.run(
        ["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", str(OCR_PS1), str(path)],
        capture_output=True, timeout=600)
    out = r.stdout.decode("utf-8", errors="replace").strip()
    if r.returncode != 0 or not out:
        err = r.stderr.decode("utf-8", errors="replace").strip()
        raise RuntimeError(err or "Windows の OCR が失敗しました")
    data = json.loads(out.splitlines()[-1])
    return {"text": data.get("text", ""), "confidence": float(data.get("confidence", 0.8)),
            "pages": 1, "method": data.get("method", "windows-ocr")}


def _read_ocr_windows(path: Path, log: Log) -> Dict:
    """Windows 版。画像はそのまま、PDF は文字層 → だめならページを画像にして OCR。"""
    if path.suffix.lower() != ".pdf":
        return _ocr_windows_image(path)

    # PDF: まず文字層を試す（画像にするより速くて正確）
    try:
        import fitz            # PyMuPDF
    except ImportError:
        raise RuntimeError("PDF を読むには PyMuPDF が要ります: pip install pymupdf")
    doc = fitz.open(str(path))
    try:
        pages = doc.page_count
        text = "\n".join(doc[i].get_text() for i in range(pages)).strip()
        if len(text) >= 20:      # 文字層があった
            return {"text": text, "confidence": 1.0, "pages": pages, "method": "pdf-text"}

        # 文字層が無い（スキャンした紙）→ ページを画像にして OCR
        log("PDF に文字が入っていないので、ページを画像にして読み取ります（%d ページ）…" % pages)
        import tempfile
        parts = []
        with tempfile.TemporaryDirectory() as tmp:
            for i in range(pages):
                png = Path(tmp) / ("p%03d.png" % i)
                doc[i].get_pixmap(dpi=200).save(str(png))
                try:
                    parts.append(_ocr_windows_image(png)["text"])
                except Exception as e:
                    log("  %d ページ目を読めませんでした: %s" % (i + 1, e))
        return {"text": "\n".join(parts), "confidence": 0.8, "pages": pages, "method": "windows-ocr-pdf"}
    finally:
        doc.close()


_whisper_model = None
_whisper_model_name = None


def _read_audio(path: Path, cfg: Dict, log: Log) -> Dict:
    if not IS_MAC:
        raise RuntimeError(
            "音声の文字起こしは macOS のみです。"
            "ほかの方法で文字起こししたテキスト（.txt）を素材フォルダに入れてください")

    global _whisper_model, _whisper_model_name
    try:
        import whisper
    except ImportError:
        raise RuntimeError("whisper が入っていません: pip3 install -U openai-whisper")
    if not shutil.which("ffmpeg"):
        raise RuntimeError("ffmpeg が見つかりません（音声の読み込みに必要）")
    name = cfg.get("model", "medium")
    if _whisper_model is None or _whisper_model_name != name:
        log("Whisper モデル「%s」を読み込んでいます（初回は少し待ちます）…" % name)
        _whisper_model = whisper.load_model(name)
        _whisper_model_name = name
    log("文字起こし中: %s（音声の長さと同じくらい時間がかかります）" % path.name)
    result = _whisper_model.transcribe(
        str(path), language=cfg.get("language", "ja"), fp16=False,
        initial_prompt=cfg.get("initialPrompt") or None, verbose=False)
    segments = result.get("segments", [])
    lines = [s["text"].strip() for s in segments if s["text"].strip()]
    duration = segments[-1]["end"] if segments else 0
    return {"text": "\n".join(lines), "duration": float(duration), "method": "whisper-" + name}


# ---------------------------------------------------------------- キャッシュ付きの入口

def _cache_key(path: Path) -> str:
    st = path.stat()
    return hashlib.sha1(("%s|%d|%d" % (path, st.st_size, st.st_mtime_ns)).encode("utf-8")).hexdigest()


def extract(path: Path, cache_dir: Path, config: Dict, log: Log) -> Dict:
    """1 ファイルから文字を取り出す。結果の dict:
        kind, text, method, confidence(OCR), duration(音声秒), pages, error
    """
    cache_dir.mkdir(parents=True, exist_ok=True)
    cache_file = cache_dir / (_cache_key(path) + ".json")
    if cache_file.exists():
        try:
            cached = json.loads(cache_file.read_text(encoding="utf-8"))
            if not cached.get("error"):
                return cached
        except Exception:
            pass

    kind = kind_of(path)
    info: Dict = {"kind": kind, "source": str(path), "text": "", "method": ""}
    try:
        if kind == "text":
            info["text"] = _read_text(path)
            info["method"] = "text"
        elif kind == "word":
            info["text"] = _read_word(path)
            if path.suffix.lower() == ".docx":
                info["method"] = "docx"
            else:
                info["method"] = "textutil" if IS_MAC else "doc97"
        elif kind == "excel":
            info["text"] = _read_excel(path)
            info["method"] = "openpyxl"
        elif kind in ("image", "pdf"):
            info.update(_read_ocr(path, log))
        elif kind == "audio":
            info.update(_read_audio(path, config.get("whisper", {}), log))
        else:
            raise RuntimeError("未対応の種類")
    except Exception as e:  # 1 件の失敗で全体を止めない
        info["error"] = str(e)
        log("  読み取り失敗: %s — %s" % (path.name, e))

    info["extractedAt"] = datetime.now().isoformat(timespec="seconds")
    try:
        cache_file.write_text(json.dumps(info, ensure_ascii=False, indent=1), encoding="utf-8")
    except Exception:
        pass
    return info


if __name__ == "__main__":
    # 動作確認用: python3 extractors.py <ファイル>
    if len(sys.argv) < 2:
        print("使い方: python3 extractors.py <ファイル>")
        sys.exit(2)
    target = Path(sys.argv[1]).expanduser().resolve()
    res = extract(target, Path(os.environ.get("TMPDIR", "/tmp")) / "nippo-cache", {}, print)
    print(json.dumps({k: v for k, v in res.items() if k != "text"}, ensure_ascii=False, indent=1))
    print(res.get("text", ""))
