"""core.py の自動テスト。

実行:  python -m unittest discover -s gikai_simple/tests
（gikai_simple フォルダの中から  python -m unittest tests.test_core  でもよい）
"""

from __future__ import annotations

import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent))

import core  # noqa: E402
import llm  # noqa: E402
import proofread  # noqa: E402

try:
    from PIL import Image
    PIL_OK = True
except ImportError:
    PIL_OK = False


def make_photo(path: Path, w: int = 1600, h: int = 1200, color=(200, 80, 80)) -> None:
    Image.new("RGB", (w, h), color).save(path, "JPEG")


class ParseTest(unittest.TestCase):
    def test_photo_line_full(self):
        ref = core.parse_photo_line("【写真】森下.jpg｜顔｜森下けい子議員")
        self.assertEqual((ref.file, ref.size, ref.caption), ("森下.jpg", "顔", "森下けい子議員"))
        self.assertEqual(ref.width_mm, core.SIZES["顔"])

    def test_photo_line_half_width_bar_and_defaults(self):
        ref = core.parse_photo_line("  【写真】 a b.png | 大 ")
        self.assertEqual((ref.file, ref.size, ref.caption), ("a b.png", "大", ""))
        ref = core.parse_photo_line("【写真】only.jpg")
        self.assertEqual((ref.file, ref.size, ref.caption), ("only.jpg", "中", ""))
        ref = core.parse_photo_line("【写真】x.jpg｜｜説明だけ")
        self.assertEqual((ref.size, ref.caption), ("中", "説明だけ"))

    def test_not_photo_line(self):
        self.assertIsNone(core.parse_photo_line("写真に兄　樹紀さん"))
        self.assertIsNone(core.parse_photo_line("【一般会計】"))

    def test_blocks_and_count(self):
        text = "見出し\n本文１\n【写真】p.jpg｜小｜説明\n本文２\n"
        blocks = core.parse_blocks(text)
        self.assertEqual([k for k, _ in blocks], ["text", "photo", "text"])
        self.assertEqual(blocks[0][1], "見出し\n本文１")
        self.assertEqual(core.count_chars(text), 9)
        self.assertEqual(core.PhotoRef("p.jpg", "小", "説明").to_line(), "【写真】p.jpg｜小｜説明")


class NumberTest(unittest.TestCase):
    def test_rules(self):
        self.assertEqual(core.normalize_numbers("4人が11月3日に第４６回"), "４人が11月３日に第46回")
        self.assertEqual(core.normalize_numbers("国道３３号、２０２５年"), "国道33号、2025年")
        self.assertEqual(core.normalize_numbers("〒７８１－２１９４　℡０８８９－２４－７７７７　本郷６１－１"),
                         "〒７８１－２１９４　℡０８８９－２４－７７７７　本郷６１－１")
        self.assertEqual(core.normalize_numbers("局 24-7777"), "局 24-7777")
        self.assertEqual(core.normalize_numbers("92・７％、１千400人"), "92・７％、１千400人")

    def test_photo_line_untouched(self):
        text = "5人\n【写真】写真1.jpg｜小｜第46回\n"
        self.assertEqual(core.normalize_numbers(text), "５人\n【写真】写真1.jpg｜小｜第46回\n")

    def test_zenkaku(self):
        self.assertEqual(core.to_zenkaku_digits("204"), "２０４")


class IssueTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def test_create_and_open(self):
        issue = core.Issue.create(self.root, "２０４", "令和８年７月31日")   # 既定は 6月号
        self.assertEqual(issue.folder.name, "第204号")
        self.assertTrue((issue.folder / "01_表紙.txt").exists())
        self.assertTrue((issue.folder / "07_裏表紙.txt").exists())
        self.assertTrue((issue.folder / "08_お知らせ.txt").exists())
        self.assertTrue((issue.folder / "別添").is_dir())
        self.assertEqual(issue.kubun_list(), core.TEMPLATES["6月号"])
        self.assertIn("第２０４号", issue.read_text("表紙"))
        again = core.Issue.open(issue.folder)
        self.assertEqual((again.gou, again.hakkoubi, again.template), ("204", "令和８年７月31日", "6月号"))
        with self.assertRaises(FileExistsError):
            core.Issue.create(self.root, "204", "")

    def test_templates_by_month(self):
        for name, extra in [("3月号", "当初予算の概要"), ("9月号", "決算の概要"), ("12月号", "行政視察・研修報告")]:
            issue = core.Issue.create(self.root, name, "", template=name)
            self.assertIn(extra, issue.kubun_list(), name)
            self.assertEqual(issue.kubun_list()[0], "表紙")
        self.assertIn("議員行政視察研修報告", core.Issue.open(self.root / "第9月号号").kubun_list())
        with self.assertRaises(KeyError):
            core.Issue.create(self.root, "x", "", template="13月号")
        for k in core.TEMPLATES.values():
            for name in k:
                self.assertIn(name, core.HINTS)   # 雛形の区分にはヒントがある

    def test_add_move_remove_kubun(self):
        issue = core.Issue.create(self.root, "1", "", template="6月号")
        issue.add_kubun("視聴者の声", after="特集")
        self.assertEqual(issue.kubun_list()[5:8], ["特集", "視聴者の声", "裏表紙"])
        self.assertTrue((issue.folder / "07_視聴者の声.txt").exists())
        self.assertTrue((issue.folder / "08_裏表紙.txt").exists())
        issue.move_kubun("視聴者の声", +1)
        self.assertEqual(issue.kubun_list()[6:8], ["裏表紙", "視聴者の声"])
        issue.move_kubun("表紙", -1)   # 端では動かない
        self.assertEqual(issue.kubun_list()[0], "表紙")
        issue.write_text("視聴者の声", "中身")
        with self.assertRaises(ValueError):
            issue.remove_kubun("視聴者の声")
        issue.write_text("視聴者の声", "")
        issue.remove_kubun("視聴者の声")
        self.assertNotIn("視聴者の声", issue.kubun_list())
        self.assertEqual([p.name for p in sorted(issue.folder.glob("0*_*.txt"))][:2], ["01_表紙.txt", "02_行政報告.txt"])
        with self.assertRaises(FileExistsError):
            issue.add_kubun("特集")
        with self.assertRaises(ValueError):
            issue.add_kubun("  ")
        # 中身は番号を付け直しても消えない
        issue.write_text("特集", "特集の本文")
        issue.add_kubun("追加", after="表紙")
        self.assertEqual(issue.read_text("特集"), "特集の本文")
        self.assertEqual(issue.kubun_list()[1], "追加")

    def test_text_roundtrip_bom(self):
        issue = core.Issue.create(self.root, "1", "")
        issue.write_text("一般質問", "質問　あ\r\n答弁　い")
        raw = issue.text_path("一般質問").read_bytes()
        self.assertTrue(raw.startswith(b"\xef\xbb\xbf"))
        self.assertEqual(issue.read_text("一般質問"), "質問　あ\n答弁　い")

    def test_read_cp932(self):
        p = self.root / "x.txt"
        p.write_bytes("議会だより".encode("cp932"))
        self.assertEqual(core.read_text_file(p), "議会だより")

    def test_import_docx_with_textbox(self):
        p = self.root / "m.docx"
        xml = (
            '<?xml version="1.0"?><w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            "<w:body><w:p><w:r><w:t>本文１</w:t></w:r></w:p>"
            "<w:p><w:r><w:t>本文</w:t></w:r><w:r><w:t>２</w:t></w:r></w:p>"
            "<w:p><w:r><w:pict><w:txbxContent><w:p><w:r><w:t>枠の中</w:t></w:r></w:p></w:txbxContent></w:pict></w:r></w:p>"
            "</w:body></w:document>"
        )
        with zipfile.ZipFile(p, "w") as z:
            z.writestr("word/document.xml", xml)
        self.assertEqual(core.import_manuscript(p), "本文１\n本文２\n枠の中")
        # 取り込み時に数字がそろう（1 桁は全角、2 桁以上は半角）
        q = self.root / "n.txt"
        q.write_text("4月と１２日", encoding="utf-8")
        self.assertEqual(core.import_manuscript(q), "４月と12日")

    def test_import_docx_rejects_doctype(self):
        p = self.root / "bad.docx"
        with zipfile.ZipFile(p, "w") as z:
            z.writestr("word/document.xml", '<!DOCTYPE x [<!ENTITY a "b">]><w:document/>')
        with self.assertRaises(ValueError):
            core.import_manuscript(p)

    def test_import_unknown_ext(self):
        p = self.root / "a.xlsx"
        p.write_bytes(b"")
        with self.assertRaises(ValueError):
            core.import_manuscript(p)


@unittest.skipUnless(PIL_OK and core.DOCX_OK, "Pillow と python-docx が必要")
class BuildTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.issue = core.Issue.create(Path(self.tmp.name), "204", "令和８年７月31日")
        make_photo(self.issue.photo_dir / "村長.jpg", 1600, 1200)
        make_photo(self.issue.photo_dir / "小さい.jpg", 300, 200)   # 解像度不足
        make_photo(self.issue.photo_dir / "余り.jpg", 800, 600)    # 使わない
        (self.issue.attach_dir / "賛否一覧.xlsx").write_bytes(b"dummy")
        self.issue.write_text("行政報告", "行政報告（要旨）\n【写真】村長.jpg｜顔｜松岡村長\n本文です。4人が第４６回に。")
        self.issue.write_text("特集", "特集の本文\n【写真】小さい.jpg｜大｜広い写真\n【写真】ない.jpg｜中")

    def tearDown(self):
        self.tmp.cleanup()

    def test_usage_and_refs(self):
        usage = self.issue.photo_usage()
        self.assertEqual(usage["村長.jpg"], ["行政報告"])
        self.assertEqual([r.file for _, r in self.issue.all_refs()], ["村長.jpg", "小さい.jpg", "ない.jpg"])

    def test_photo_info(self):
        info = core.photo_info(self.issue.photo_dir / "小さい.jpg")
        self.assertEqual((info.width_px, info.height_px), (300, 200))
        self.assertIn("解像度不足", info.warning_for(80))
        self.assertEqual(info.warning_for(20), "")
        self.assertIn("ありません", core.photo_info(self.issue.photo_dir / "ない.jpg").error)

    def test_build_all(self):
        outs, warnings = core.build_all(self.issue)
        self.assertEqual(len(outs), 2)
        for p in outs:
            self.assertTrue(p.exists() and p.stat().st_size > 0, p)
        joined = "\n".join(warnings)
        self.assertIn("ない.jpg", joined)
        self.assertIn("解像度不足", joined)
        self.assertNotIn("表紙: 原稿が空です", warnings)   # 表紙は雛形入りなので空ではない
        self.assertIn("一般質問: 原稿が空です", warnings)

        from docx import Document
        doc = Document(outs[0])
        texts = [p.text for p in doc.paragraphs]
        self.assertTrue(any(t.startswith("■ 行政報告") for t in texts))
        self.assertIn("本文です。４人が第46回に。", texts)
        self.assertIn("第２０４号", texts)   # 表紙は変換しない
        self.assertTrue(any(t.startswith("【写真1】村長.jpg（顔・幅26mm）") and "松岡村長" in t for t in texts))
        self.assertTrue(any(t.startswith("【写真3】ない.jpg") for t in texts))
        self.assertEqual(len(doc.inline_shapes), 2)   # 見つかった写真 2 枚だけ貼られる
        for shp in doc.inline_shapes:
            self.assertLessEqual(shp.height.mm, core.DAN_HEIGHT_MM)   # 1 段の高さに収まる
        # 表紙のセクションは横書き 1 段、本文のセクションは縦書き 5 段。指示書は横書き
        from docx.oxml.ns import qn
        self.assertEqual(len(doc.sections), 2)
        cover, body = doc.sections
        self.assertIsNone(cover._sectPr.find(qn("w:textDirection")))
        td = body._sectPr.find(qn("w:textDirection"))
        self.assertIsNotNone(td)
        self.assertEqual(td.get(qn("w:val")), "tbRl")
        self.assertEqual(list(body._sectPr).index(td) + 1,
                         list(body._sectPr).index(body._sectPr.find(qn("w:docGrid"))))
        cols = body._sectPr.find(qn("w:cols"))
        self.assertEqual(cols.get(qn("w:num")), "5")
        self.assertEqual(cols.get(qn("w:space")), "340")
        self.assertEqual(round(body.top_margin.mm), 15)
        # 2 桁の数字には縦中横が掛かる
        para = next(p for p in doc.paragraphs if p.text == "本文です。４人が第46回に。")
        runs = [(r.text, r._element.find(".//" + qn("w:eastAsianLayout")) is not None) for r in para.runs]
        self.assertEqual(runs, [("本文です。４人が第", False), ("46", True), ("回に。", False)])
        # 表紙には掛からない
        cpara = next(p for p in doc.paragraphs if p.text.startswith("令和"))
        self.assertTrue(all(r._element.find(".//" + qn("w:eastAsianLayout")) is None for r in cpara.runs))

        sheet = Document(outs[1])
        self.assertIsNone(sheet.sections[0]._sectPr.find(qn("w:textDirection")))
        table = sheet.tables[0]
        self.assertEqual(len(table.rows), 1 + 3)
        self.assertEqual(table.rows[1].cells[2].paragraphs[0].text, "村長.jpg")
        self.assertIn("余り.jpg", "\n".join(p.text for p in sheet.paragraphs))
        self.assertIn("・賛否一覧.xlsx", [p.text for p in sheet.paragraphs])   # 別添の一覧


class WidthTest(unittest.TestCase):
    """組版で何文字ぶんになるか（半角は 0.5）。"""

    def test_count_width(self):
        self.assertEqual(core.count_width("あいう"), 3)
        self.assertEqual(core.count_width("12日"), 2)        # 0.5+0.5+1
        self.assertEqual(core.count_width("１２日"), 3)
        self.assertEqual(core.count_width("あ い\nう"), 3)    # 空白と改行は数えない
        self.assertEqual(core.count_width("あ\n【写真】x.jpg｜大｜説明"), 1)   # 写真行は数えない

    def test_page_capacity(self):
        cap = core.page_capacity()
        # 第203号の紙面（5段・10.5pt・行送り1.45）で 13字 × 33行 × 5段 = 2145字
        self.assertEqual(cap["chars_per_line"], 13)
        self.assertEqual(cap["lines_per_dan"], 33)
        self.assertEqual(cap["chars_per_page"], 2145)
        self.assertAlmostEqual(cap["char_area_mm2"], 19.9, delta=0.2)

    def test_photo_chars(self):
        small = core.photo_chars(core.PhotoRef("a.jpg", "顔"))
        big = core.photo_chars(core.PhotoRef("a.jpg", "大"))
        self.assertLess(small, big)                     # 大きい写真ほど多く押しのける
        self.assertGreater(small, 0)
        # キャプションがあるぶんだけ増える
        self.assertGreater(core.photo_chars(core.PhotoRef("a.jpg", "中", "説明")),
                           core.photo_chars(core.PhotoRef("a.jpg", "中")))


class ProofreadTest(unittest.TestCase):
    """ルールベースの校正。辞書は proof_data/ の JSON。"""

    def find(self, text, category):
        return [i for i in proofread.proofread(text) if i.category == category]

    def test_style_and_typo(self):
        self.assertTrue(self.find("出来るだけ早く", "kana"))
        self.assertTrue(self.find("課題等について", "house"))
        self.assertTrue(self.find("①番目", "symbols"))
        self.assertTrue(self.find("まず最初に", "redundant"))

    def test_grammar_and_punct(self):
        self.assertTrue(self.find("これは見れると思う", "ら抜き言葉"))
        self.assertTrue(self.find("（かっこが閉じない。", "括弧"))

    def test_proper_noun(self):
        got = self.find("山崎副村長から説明があった。", "固有名詞")
        self.assertTrue(got)
        self.assertEqual(got[0].suggestion, "山﨑副村長")

    def test_no_false_positive_on_clean_text(self):
        clean = "１月15日に第１回臨時会が開催され、議案１件が可決された。"
        self.assertEqual([i for i in proofread.proofread(clean) if i.severity == "error"], [])

    def test_apply_fixes(self):
        text = "出来るだけ早く、山崎副村長が対応する。"
        fixed = proofread.apply_fixes(text, proofread.proofread(text))
        self.assertIn("できるだけ", fixed)
        self.assertIn("山﨑副村長", fixed)

    def test_ambiguous_fix_not_applied(self):
        # 「まず／最初に」のようにどちらか選ぶものは、勝手に直さない
        text = "まず最初に検討する。"
        self.assertEqual(proofread.apply_fixes(text, proofread.proofread(text)), text)


class LlmTest(unittest.TestCase):
    """LLM の差し込み口。Ollama が無くても壊れないこと。"""

    def test_host_is_locked_to_localhost(self):
        for bad in ("example.com", "10.0.0.1", "api.openai.com"):
            with self.assertRaises(ValueError):
                llm._check_host(bad)
        for good in llm.ALLOWED_HOSTS:
            llm._check_host(good)      # 例外が出ないこと

    def test_available_without_ollama(self):
        ok, msg = llm.available()
        if not ok:                      # Ollama が入っていない普通の環境
            self.assertIn("AI 校正", msg)

    def test_chunks_split_at_sentence_end(self):
        text = "一文目です。" * 200
        parts = list(llm._chunks(text))
        self.assertGreater(len(parts), 1)
        for off, c in parts:
            self.assertEqual(text[off:off + len(c)], c)      # 位置がずれていない
            self.assertLessEqual(len(c), llm.CHUNK_CHARS + 20)

    def test_pick_model_prefers_larger(self):
        self.assertEqual(llm.pick_model(["qwen2.5:3b", "qwen2.5:7b"]), "qwen2.5:7b")
        self.assertEqual(llm.pick_model(["qwen2.5:3b"]), "qwen2.5:3b")
        self.assertEqual(llm.pick_model(["mystery:1b"]), "mystery:1b")   # 知らないものでも使う
        self.assertEqual(llm.pick_model([]), "")

    def test_is_noise(self):
        nouns = {"武政義幸", "日高村", "山﨑副村長"}
        # 直っていない指摘（実際に 3b が返してきた形）
        self.assertTrue(llm._is_noise("武政義幸様", "武政義幸様", nouns))
        self.assertTrue(llm._is_noise("大祭", "", nouns))
        # 固有名詞には触らせない（辞書の担当）
        self.assertTrue(llm._is_noise("日高村", "日高町", nouns))
        self.assertTrue(llm._is_noise("山﨑副村長が述べた", "山崎副村長が述べた", nouns))
        # 文まるごとの書き直し
        self.assertTrue(llm._is_noise("あ" * 50, "い" * 50, nouns))
        # まるで別物への言い換え
        self.assertTrue(llm._is_noise("検討する", "前向きに善処してまいります", nouns))
        # 本当の誤字の直しは通す
        self.assertFalse(llm._is_noise("説明ました", "説明しました", nouns))
        self.assertFalse(llm._is_noise("実施ます", "実施します", nouns))

    def test_diff_ratio(self):
        self.assertLess(llm._diff_ratio("説明ました", "説明しました"), 0.2)
        self.assertGreater(llm._diff_ratio("検討する", "まったく別の文章です"), 0.5)

    def test_parse_loose_json(self):
        got = llm._parse('はい。\n[{"text":"実施ます","fix":"実施します","why":"脱字"}]\nご確認ください')
        self.assertEqual(got[0]["fix"], "実施します")
        self.assertEqual(llm._parse("見つかりませんでした"), [])
        self.assertEqual(llm._parse("[壊れたJSON"), [])


class EstimateTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.issue = core.Issue.create(Path(self.tmp.name), "204", "", template="6月号")

    def tearDown(self):
        self.tmp.cleanup()

    def test_estimate(self):
        self.issue.write_text("一般質問", "あ" * 2145)      # ちょうど 1 ページぶん
        est = self.issue.estimate()
        row = next(r for r in est["rows"] if r["kubun"] == "一般質問")
        self.assertEqual(row["chars"], 2145)
        self.assertAlmostEqual(row["pages"], 1.0, places=2)
        self.assertTrue(next(r for r in est["rows"] if r["kubun"] == "表紙")["cover"])
        # 表紙 1 ページ + 本文 1 ページ = 2 ページ（偶数）
        self.assertEqual(est["pages"], 2)
        self.assertFalse(est["odd"])

    def test_odd_pages_flagged(self):
        self.issue.write_text("一般質問", "あ" * 2145 * 2)
        est = self.issue.estimate()
        self.assertEqual(est["pages"], 3)
        self.assertTrue(est["odd"])


if __name__ == "__main__":
    unittest.main()
